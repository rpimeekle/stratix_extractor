"""Load the list of switches to connect to from a CSV or Excel file.

Expected columns (header names are case-insensitive, extra columns ignored):

    hostname, ip, username, password, enable_secret, port

Only the IP address is mandatory. Blank usernames/passwords fall back to the
default credentials entered in the GUI.
"""

from __future__ import annotations

import csv
import ipaddress
from dataclasses import dataclass, field
from pathlib import Path

TEMPLATE_HEADERS = ["hostname", "ip", "username", "password", "enable_secret", "port"]

_ALIASES = {
    "hostname": {"hostname", "host name", "name", "switch", "switch name", "device", "device name"},
    "ip": {"ip", "ip address", "ip_address", "ipaddress", "address", "host", "management ip", "mgmt ip"},
    "username": {"username", "user name", "user", "login"},
    "password": {"password", "pass", "pwd"},
    "enable_secret": {"enable_secret", "enable secret", "enable", "secret", "enable password"},
    "port": {"port", "ssh port", "ssh_port"},
}


@dataclass
class SwitchTarget:
    hostname: str
    ip: str
    username: str = ""
    password: str = field(default="", repr=False)
    enable_secret: str = field(default="", repr=False)
    port: int = 22
    row: int = 0  # row number in the source file, for messages

    @property
    def label(self) -> str:
        return f"{self.hostname} ({self.ip})" if self.hostname != self.ip else self.ip

    @property
    def has_file_credentials(self) -> bool:
        return bool(self.username and self.password)


@dataclass
class Credentials:
    username: str = ""
    password: str = field(default="", repr=False)
    enable_secret: str = field(default="", repr=False)


def resolve_credentials(target: SwitchTarget, defaults: Credentials, override: bool) -> Credentials:
    """Credentials to use for one switch.

    With override on, the GUI credentials are used for every switch. Otherwise
    each field comes from the file and blanks are filled from the defaults.
    """
    if override:
        return Credentials(
            defaults.username,
            defaults.password,
            defaults.enable_secret or target.enable_secret,
        )
    return Credentials(
        target.username or defaults.username,
        target.password or defaults.password,
        target.enable_secret or defaults.enable_secret,
    )


def _canonical(header: str) -> str | None:
    key = (header or "").strip().lower().replace("-", " ").replace("_", " ")
    key_us = key.replace(" ", "_")
    for canonical, names in _ALIASES.items():
        normalised = {n.replace("_", " ") for n in names}
        if key in normalised or key_us in names:
            return canonical
    return None


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _read_rows(path: Path) -> list[list[str]]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.worksheets[0]
            return [[_clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
        finally:
            workbook.close()
    if suffix in (".csv", ".txt", ""):
        # utf-8-sig removes the BOM Excel adds when saving CSV.
        with open(path, newline="", encoding="utf-8-sig") as handle:
            sample = handle.read(4096)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [[_clean(cell) for cell in row] for row in csv.reader(handle, dialect)]
    raise ValueError(f"Unsupported file type '{path.suffix}'. Use .csv or .xlsx.")


def load_switch_list(path: str | Path) -> tuple[list[SwitchTarget], list[str]]:
    """Return (targets, warnings). Raises ValueError if the file is unusable."""
    path = Path(path)
    rows = _read_rows(path)
    warnings: list[str] = []

    header_index = None
    columns: dict[str, int] = {}
    for index, row in enumerate(rows):
        if not any(row) or (row[0].startswith("#")):
            continue
        mapping = {}
        for col, cell in enumerate(row):
            canonical = _canonical(cell)
            if canonical and canonical not in mapping:
                mapping[canonical] = col
        if "ip" in mapping:
            header_index, columns = index, mapping
        break
    if header_index is None:
        raise ValueError(
            "Could not find a header row with an 'ip' column. "
            "Expected columns: " + ", ".join(TEMPLATE_HEADERS)
        )

    def cell(row: list[str], name: str) -> str:
        col = columns.get(name)
        return row[col].strip() if col is not None and col < len(row) else ""

    targets: list[SwitchTarget] = []
    seen_ips: set[str] = set()
    for index in range(header_index + 1, len(rows)):
        row = rows[index]
        row_number = index + 1
        if not any(row) or row[0].startswith("#"):
            continue
        ip = cell(row, "ip")
        hostname = cell(row, "hostname")
        if not ip:
            warnings.append(f"Row {row_number}: no IP address, skipped ({hostname or 'unnamed'}).")
            continue
        try:
            ipaddress.ip_address(ip)
        except ValueError:
            # DNS names are allowed; just flag anything with spaces.
            if " " in ip:
                warnings.append(f"Row {row_number}: '{ip}' is not a valid address, skipped.")
                continue
        if ip in seen_ips:
            warnings.append(f"Row {row_number}: {ip} is listed more than once.")
        seen_ips.add(ip)

        port_text = cell(row, "port")
        port = 22
        if port_text:
            try:
                port = int(float(port_text))
            except ValueError:
                warnings.append(f"Row {row_number}: SSH port '{port_text}' is not a number, using 22.")

        targets.append(SwitchTarget(
            hostname=hostname or ip,
            ip=ip,
            username=cell(row, "username"),
            password=cell(row, "password"),
            enable_secret=cell(row, "enable_secret"),
            port=port,
            row=row_number,
        ))
    if not targets:
        raise ValueError("The file has a header row but no switches with an IP address.")
    return targets, warnings


def write_template(path: str | Path) -> None:
    """Write an example switch list as CSV or XLSX depending on the extension."""
    path = Path(path)
    example = [
        ["LINE3-SW01", "10.20.3.11", "", "", "", "22"],
        ["PKG-SW52", "10.20.4.52", "admin", "ChangeMe", "", "22"],
    ]
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Switches"
        sheet.append(TEMPLATE_HEADERS)
        for row in example:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col, width in zip("ABCDEF", (18, 16, 14, 14, 16, 8)):
            sheet.column_dimensions[col].width = width
        workbook.save(path)
    else:
        with open(path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(TEMPLATE_HEADERS)
            writer.writerows(example)
