import tempfile
import threading
from pathlib import Path

from openpyxl import load_workbook

from stratix_extractor.collector import Collector
from stratix_extractor.excel_writer import check_writable
from stratix_extractor.runner import run_extraction
from stratix_extractor.switchlist import (Credentials, SwitchTarget, load_switch_list,
                                          resolve_credentials, write_template)

from .fakes import NetmikoAuthenticationException, make_connect


def _targets():
    return [
        SwitchTarget("LINE3-SW01", "10.20.3.11"),
        SwitchTarget("PKG-SW52", "10.20.4.52", username="admin", password="pw"),
        SwitchTarget("BAD-SW", "10.20.9.9"),
    ]


def _run(tmp: Path, stop_event=None):
    connect = make_connect({
        "10.20.3.11": ("stratix5700", "LINE3-SW01#"),
        "10.20.4.52": ("stratix5200", "PKG-SW52>"),
        "10.20.9.9": NetmikoAuthenticationException("Authentication to device failed."),
    })
    log = []
    collector = Collector(lambda m, lvl: log.append((lvl, m)), raw_dir=tmp / "raw", connect=connect)
    report = run_extraction(
        _targets(), Credentials("default", "defpw"), False, tmp / "out.xlsx", "switches.csv",
        lambda m, lvl: log.append((lvl, m)), collector, stop_event=stop_event,
    )
    return report, log, connect


def test_end_to_end_workbook():
    with tempfile.TemporaryDirectory() as folder:
        tmp = Path(folder)
        report, log, connect = _run(tmp)
        assert [r.success for r in report.results] == [True, True, False]
        assert "Authentication failed" in report.results[2].error
        assert report.verified, report.verify_message
        assert all(conn.disconnected for conn in connect.made.values())
        # 5200 logged in at '>' with no enable secret -> warning, still collected
        assert any("user EXEC" in w for w in report.results[1].warnings)

        wb = load_workbook(report.output_path)
        assert wb.sheetnames == ["Summary", "Switches", "Ports", "Inventory", "VLANs", "Neighbours"]
        ports = wb["Ports"]
        headers = [c.value for c in ports[1]]
        rows = [dict(zip(headers, r)) for r in ports.iter_rows(min_row=2, values_only=True)]
        assert len(rows) == 10 + 5
        fa16 = next(r for r in rows if r["Switch"] == "LINE3-SW01" and r["Port"] == "FastEthernet1/6")
        assert fa16["Description"] == "Camera CAM-12 vlan 30 - packaging hall north wall, PoE injector PI-4"
        assert fa16["Status"] == "connected" and fa16["Access VLAN"] == "30 (CAMERAS)"
        fa11 = next(r for r in rows if r["Port"] == "FastEthernet1/1")
        assert fa11["Neighbour"] == "L3_FILLER_PLC" and fa11["Neighbour IP"] == "10.20.3.50"
        gi11 = next(r for r in rows if r["Switch"] == "LINE3-SW01" and r["Port"] == "GigabitEthernet1/1")
        assert gi11["Mode"] == "trunk" and gi11["Allowed VLANs"].startswith("1,10,20")

        switches = wb["Switches"]
        sw_headers = [c.value for c in switches[1]]
        sw = [dict(zip(sw_headers, r)) for r in switches.iter_rows(min_row=2, values_only=True)]
        assert [s["Family"] for s in sw] == ["Stratix 5700", "Stratix 5200"]
        assert sw[0]["Ports err-disabled"] == 1
        assert sw[1]["Management IPs"] == "Vlan1 10.20.4.52"

        summary = [r for r in wb["Summary"].iter_rows(values_only=True) if r and r[2] in ("Success", "Failed")]
        assert [r[2] for r in summary] == ["Success", "Success", "Failed"]
        assert (tmp / "raw" / "LINE3-SW01_10.20.3.11" / "show_interfaces.txt").exists()
        wb.close()


def test_stop_before_start_writes_nothing():
    with tempfile.TemporaryDirectory() as folder:
        stop = threading.Event()
        stop.set()
        report, _, _ = _run(Path(folder), stop_event=stop)
        assert report.stopped and report.results == [] and report.output_path is None


def test_switch_list_csv_and_xlsx_with_aliases():
    with tempfile.TemporaryDirectory() as folder:
        tmp = Path(folder)
        csv_file = tmp / "list.csv"
        csv_file.write_text(
            "﻿Host Name,IP Address,User,Password,Port\n"
            "SW1,10.0.0.1,,,\n"
            "# commented out,10.0.0.9,,,\n"
            ",10.0.0.2,bob,secret,2222\n"
            "SW3,,x,y,\n"
            "SW4,10.0.0.1,,,\n",
            encoding="utf-8",
        )
        targets, warnings = load_switch_list(csv_file)
        assert [(t.hostname, t.ip, t.port) for t in targets] == [
            ("SW1", "10.0.0.1", 22), ("10.0.0.2", "10.0.0.2", 2222), ("SW4", "10.0.0.1", 22)]
        assert any("no IP address" in w for w in warnings)
        assert any("more than once" in w for w in warnings)

        xlsx = tmp / "template.xlsx"
        write_template(xlsx)
        targets, warnings = load_switch_list(xlsx)
        assert [t.ip for t in targets] == ["10.20.3.11", "10.20.4.52"] and not warnings


def test_credential_resolution():
    t = SwitchTarget("A", "1.1.1.1", username="fileuser", password="")
    defaults = Credentials("gui", "guipw", "en")
    assert resolve_credentials(t, defaults, False) == Credentials("fileuser", "guipw", "en")
    assert resolve_credentials(t, defaults, True) == Credentials("gui", "guipw", "en")


def test_check_writable():
    with tempfile.TemporaryDirectory() as folder:
        assert check_writable(Path(folder) / "new.xlsx") is None
        assert not (Path(folder) / "new.xlsx").exists()
        assert "does not exist" in check_writable(Path(folder) / "missing" / "x.xlsx")


def _netmiko_style_failure(os_error):
    """Raise the way Netmiko does: its own exception, raised inside 'except socket.error'."""
    from .fakes import NetmikoTimeoutException

    def connect(params):
        try:
            raise os_error
        except OSError:
            raise NetmikoTimeoutException(  # noqa: B904 - mimics Netmiko
                "TCP connection to device failed.\nCommon causes of this problem are:\n"
                "1. Incorrect hostname or IP address.\n2. Wrong TCP port.\n"
                "3. Intermediate firewall blocking access.")
    return connect


def test_immediate_connection_failure_reports_real_cause_and_probes_ports():
    log = []
    probed = []

    def probe(host, ports):
        probed.append((host, ports))
        return {22: "connection refused [error 10061]", 23: "open"}

    collector = Collector(lambda m, lvl: log.append((lvl, m)),
                          connect=_netmiko_style_failure(ConnectionRefusedError(10061, "refused")),
                          probe=probe)
    result = collector.collect(SwitchTarget("SW", "172.21.117.11"), Credentials("admin", "pw"))
    assert not result.success
    assert "connection refused" in result.error and "within" not in result.error
    assert "[error 10061]" in result.error
    assert probed == [("172.21.117.11", (22, 23))]
    text = "\n".join(m for _, m in log)
    assert "Port 23 (Telnet): open" in text
    assert "SSH is probably not enabled" in text


def test_blocked_on_this_pc_hint():
    log = []
    blocked = "blocked on this PC [error 10013]"
    collector = Collector(lambda m, lvl: log.append(m),
                          connect=_netmiko_style_failure(PermissionError(10013, "forbidden")),
                          probe=lambda host, ports: {p: blocked for p in ports})
    result = collector.collect(SwitchTarget("SW", "10.0.0.1"), Credentials("a", "b"))
    assert "firewall or endpoint-protection" in result.error
    assert any("Ask IT to allow" in m for m in log)


def test_genuine_timeout_still_says_timeout():
    from stratix_extractor.collector import _describe_error

    from .fakes import NetmikoTimeoutException
    target = SwitchTarget("SW", "10.0.0.1")
    exc = NetmikoTimeoutException("Connection to device timed-out")
    assert "within 20 s" in _describe_error(exc, target, 20, elapsed=20.3)
    assert "within" not in _describe_error(exc, target, 20, elapsed=0.02)


def test_probe_ports_against_closed_local_port():
    import socket

    from stratix_extractor.collector import probe_ports
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        port = s.getsockname()[1]  # bound but not listening -> refused once closed
    # Windows retries a refused SYN for ~2 s before reporting it, so allow longer.
    result = probe_ports("127.0.0.1", (port,), timeout=10)
    assert "refused" in result[port]


def test_incompatible_peer_explained():
    class IncompatiblePeer(Exception):
        pass

    def connect(params):
        try:
            raise IncompatiblePeer("Incompatible ssh peer (no acceptable kex algorithm)")
        except IncompatiblePeer as err:
            from .fakes import NetmikoTimeoutException
            raise NetmikoTimeoutException(
                "A paramiko SSHException occurred during connection creation: "
                "Incompatible ssh peer (no acceptable kex algorithm)") from err

    log = []
    collector = Collector(lambda m, lvl: log.append(m), connect=connect,
                          probe=lambda h, p: {})
    result = collector.collect(SwitchTarget("SW", "10.0.0.1"), Credentials("a", "b"))
    assert "no SSH algorithm in common" in result.error
    assert "no acceptable kex algorithm" in result.error
    assert not any("Checking TCP ports" in m for m in log)


def test_bundled_paramiko_supports_stratix_legacy_algorithms():
    """IOS 15.x Stratix switches need SHA-1 key exchange and ssh-rsa host keys.
    Paramiko 5 removed them; this fails the build if they disappear again."""
    import pytest

    pytest.importorskip("paramiko")
    from stratix_extractor.collector import missing_legacy_support

    assert missing_legacy_support() == []
